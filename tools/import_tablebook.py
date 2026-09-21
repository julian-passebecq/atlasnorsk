#!/usr/bin/env python3
"""Convert an XLSX/XLSM workbook into AtlasNorsk's native Table Book JSON.

This is an offline import utility, not an application runtime dependency.
It deliberately preserves values and tabs, not Excel formatting/macros/formulas.

Usage:
    python tools/import_tablebook.py source.xlsx output.json --id tablebook.my-book --title "My Book"
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from openpyxl import load_workbook


def slug(value: str) -> str:
    value = value.lower().replace("å", "a").replace("ø", "o").replace("æ", "ae")
    return re.sub(r"[^a-z0-9]+", "-", value).strip("-") or "sheet"


def clean(value):
    if value is None:
        return None
    if isinstance(value, (int, float, bool)):
        return value
    return str(value).replace("\r\n", "\n").strip()


def parse_sheet(ws):
    raw = []
    for row in ws.iter_rows(values_only=True):
        values = [clean(value) for value in row]
        while values and values[-1] is None:
            values.pop()
        if any(value not in (None, "") for value in values):
            raw.append(values)

    if not raw:
        return {"columns": [], "preface": [], "rows": [], "rowCount": 0}

    header_index = 0
    best_score = -1
    likely_headers = {
        "norsk", "english", "french", "type", "category", "forms", "pattern",
        "gender", "source", "chapter", "marker", "forklaring", "synonyms", "antonyms",
    }

    for index, row in enumerate(raw[:12]):
        populated = [value for value in row if value not in (None, "")]
        words = {str(value).strip().lower() for value in populated}
        score = len(populated) + 5 * len(words & likely_headers)
        if len(populated) >= 3 and score > best_score:
            best_score = score
            header_index = index

    header = [
        str(value) if value not in (None, "") else f"Column {index + 1}"
        for index, value in enumerate(raw[header_index])
    ]
    width = len(header)
    rows = []

    for index, row in enumerate(raw):
        if index == header_index:
            continue
        cells = (row + [None] * max(0, width - len(row)))[:width]
        non_empty = sum(value not in (None, "") for value in cells)
        kind = "section" if non_empty == 1 else ("meta" if index < header_index else "data")
        rows.append({"kind": kind, "cells": cells})

    return {
        "columns": header,
        "preface": raw[:header_index],
        "rows": rows,
        "rowCount": len(rows),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source")
    parser.add_argument("output")
    parser.add_argument("--id", required=True)
    parser.add_argument("--title", required=True)
    parser.add_argument("--level", default="Unknown")
    parser.add_argument("--sheets", nargs="*")
    args = parser.parse_args()

    source = Path(args.source)
    workbook = load_workbook(
        source,
        read_only=True,
        data_only=True,
        keep_vba=source.suffix.lower() == ".xlsm",
    )
    names = args.sheets or workbook.sheetnames

    tabs = []
    for name in names:
        parsed = parse_sheet(workbook[name])
        tabs.append({
            "id": slug(name),
            "title": name,
            "sourceSheet": name,
            **parsed,
        })
    workbook.close()

    payload = {
        "schemaVersion": 1,
        "id": args.id,
        "title": args.title,
        "kind": "tablebook",
        "level": args.level,
        "sourceFile": source.name,
        "tabs": tabs,
    }

    Path(args.output).write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
